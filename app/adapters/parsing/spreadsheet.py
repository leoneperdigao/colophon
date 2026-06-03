"""Spreadsheet DocumentParser adapter (openpyxl).

Reads cell **values, never formulas** (`data_only=True`), bounded by a sheet cap.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.application.ports.document_parser import ParseError
from app.domain.stage_result import Curated


class SpreadsheetParser:
    def __init__(self, *, max_sheets: int = 20) -> None:
        self._max_sheets = max_sheets

    def parse(self, content: bytes, content_type: str, filename: str) -> Curated:
        try:
            workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
        except Exception as exc:  # malformed/abusive input -> stable message; detail stays in cause
            raise ParseError("could not read the spreadsheet document") from exc
        try:
            sheet_names = list(workbook.sheetnames)
            if len(sheet_names) > self._max_sheets:
                raise ParseError(f"spreadsheet exceeds the sheet limit ({len(sheet_names)} sheets)")
            lines: list[str] = []
            for name in sheet_names:
                for row in workbook[name].iter_rows(values_only=True):
                    cells = [str(value) for value in row if value is not None]
                    if cells:
                        lines.append("\t".join(cells))
        finally:
            workbook.close()
        return Curated(
            text="\n".join(lines),
            page_or_sheet_count=len(sheet_names),
            structure={"sheets": sheet_names},
        )
