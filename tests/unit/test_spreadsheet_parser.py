"""SpreadsheetParser adapter tests — real openpyxl, values not formulas. T042."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.adapters.parsing.spreadsheet import SpreadsheetParser
from app.application.ports.document_parser import ParseError

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_xlsx(sheets: dict[str, list[tuple[str, object]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(title=name)
        for ref, value in cells:
            ws[ref] = value
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_reads_values_and_never_leaks_formulas() -> None:
    content = _make_xlsx({"Data": [("A1", "Hello"), ("A2", "World"), ("B1", "=1+2")]})
    curated = SpreadsheetParser().parse(content, XLSX, "f.xlsx")
    assert "Hello" in curated.text
    assert "World" in curated.text
    assert "=1+2" not in curated.text  # formula string must never appear
    assert curated.page_or_sheet_count == 1


def test_counts_sheets_and_enforces_cap() -> None:
    content = _make_xlsx({"S1": [("A1", "x")], "S2": [("A1", "y")]})
    assert SpreadsheetParser().parse(content, XLSX, "f.xlsx").page_or_sheet_count == 2
    with pytest.raises(ParseError):
        SpreadsheetParser(max_sheets=1).parse(content, XLSX, "f.xlsx")


def test_corrupt_spreadsheet_raises_user_safe_parse_error() -> None:
    with pytest.raises(ParseError) as exc_info:
        SpreadsheetParser().parse(b"not a spreadsheet", XLSX, "SECRET-NAME.xlsx")
    message = str(exc_info.value)
    assert message == "could not read the spreadsheet document"  # stable, no internals
    assert "SECRET-NAME" not in message  # client-supplied filename not echoed
    assert exc_info.value.__cause__ is not None  # detail preserved as chained cause (logs only)
