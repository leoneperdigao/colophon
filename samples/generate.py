"""Generate the eval **gold set** — documents synthesised from known structured
records, so the expected document_type and key_entities are known by construction
(free, exact labels). Distinct from the downloaded, unlabeled robustness corpus.

Every entity value is written verbatim into the document text, so a correct
extraction is also grounded.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from fpdf import FPDF
from openpyxl import Workbook

from app.adapters.parsing.content_types import PDF, XLSX


@dataclass(frozen=True)
class GoldSample:
    content: bytes
    content_type: str
    filename: str
    document_type: str
    key_entities: tuple[tuple[str, str], ...]  # (type, value)


def _pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def _xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _xlsx_sheets(sheets: list[tuple[str, list[list[str]]]]) -> bytes:
    """Build a workbook with explicitly named sheets (multi-sheet coverage)."""
    workbook = Workbook()
    workbook.remove(workbook.active)  # drop the default sheet; name every sheet ourselves
    for title, rows in sheets:
        sheet = workbook.create_sheet(title=title)
        for row in rows:
            sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_gold_set() -> list[GoldSample]:
    return [
        GoldSample(
            content=_pdf(
                "INVOICE\nFrom: ACME Corporation\nInvoice date: 2026-01-15\nAmount due: USD 4200\n"
            ),
            content_type=PDF,
            filename="gold-invoice-001.pdf",
            document_type="invoice",
            key_entities=(
                ("org", "ACME Corporation"),
                ("date", "2026-01-15"),
                ("amount", "USD 4200"),
            ),
        ),
        GoldSample(
            content=_pdf(
                "Quarterly Report\n"
                "Prepared by Globex Industries\n"
                "Reporting period ending 2026-03-31\n"
            ),
            content_type=PDF,
            filename="gold-report-001.pdf",
            document_type="report",
            key_entities=(
                ("org", "Globex Industries"),
                ("date", "2026-03-31"),
            ),
        ),
        GoldSample(
            content=_pdf(
                "Dear Initech,\nThis letter confirms our meeting on 2026-05-09.\nRegards, Hooli\n"
            ),
            content_type=PDF,
            filename="gold-letter-001.pdf",
            document_type="letter",
            key_entities=(
                ("org", "Initech"),
                ("org", "Hooli"),
                ("date", "2026-05-09"),
            ),
        ),
        GoldSample(
            content=_xlsx(
                [
                    ["Vendor", "Umbrella Corp"],
                    ["Total", "USD 1500"],
                    ["Date", "2026-02-02"],
                ]
            ),
            content_type=XLSX,
            filename="gold-sheet-001.xlsx",
            document_type="spreadsheet",
            key_entities=(
                ("org", "Umbrella Corp"),
                ("amount", "USD 1500"),
                ("date", "2026-02-02"),
            ),
        ),
        # A second invoice — denser text (an invoice number that is deliberately
        # NOT a gold entity), so extraction must pick the labelled fields out of noise.
        GoldSample(
            content=_pdf(
                "INVOICE\n"
                "From: Stark Industries\n"
                "Invoice number: INV-2026-0042\n"
                "Invoice date: 2026-04-20\n"
                "Amount due: USD 9875\n"
            ),
            content_type=PDF,
            filename="gold-invoice-002.pdf",
            document_type="invoice",
            key_entities=(
                ("org", "Stark Industries"),
                ("date", "2026-04-20"),
                ("amount", "USD 9875"),
            ),
        ),
        # A second report with two distinct organisations — tests multi-entity
        # extraction of the same type (org) from one document.
        GoldSample(
            content=_pdf(
                "Annual Financial Report\n"
                "Prepared by Wayne Enterprises\n"
                "Audited by Daily Planet Auditors\n"
                "Fiscal year ending 2026-12-31\n"
                "Net revenue: USD 12000\n"
            ),
            content_type=PDF,
            filename="gold-report-002.pdf",
            document_type="report",
            key_entities=(
                ("org", "Wayne Enterprises"),
                ("org", "Daily Planet Auditors"),
                ("date", "2026-12-31"),
                ("amount", "USD 12000"),
            ),
        ),
        # A multi-sheet workbook — entities spread across sheets, so a correct
        # extraction depends on the parser concatenating every sheet (sheet-cap path).
        GoldSample(
            content=_xlsx_sheets(
                [
                    ("Summary", [["Client", "Cyberdyne Systems"], ["Period", "2026-07-15"]]),
                    ("Transactions", [["Item", "Amount"], ["Subscription", "USD 7300"]]),
                ]
            ),
            content_type=XLSX,
            filename="gold-sheet-002.xlsx",
            document_type="spreadsheet",
            key_entities=(
                ("org", "Cyberdyne Systems"),
                ("date", "2026-07-15"),
                ("amount", "USD 7300"),
            ),
        ),
    ]
